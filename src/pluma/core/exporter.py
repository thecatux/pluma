"""Exportación del texto de una novela a distintos formatos.

Construye un único QTextDocument (portada opcional + capítulos) y lo
vuelca al formato elegido. Qt sabe escribir PDF, ODT, Markdown y texto
plano por sí mismo; solo .docx necesita `python-docx`, que es opcional
y únicamente se exige al exportar a ese formato. Esta capa no tiene
interfaz: el panel "Exportar" solo elige capítulos, formato y ruta.
"""

from html import escape as html_escape
from pathlib import Path

from PySide6.QtCore import Qt, QByteArray, QMarginsF
from PySide6.QtGui import (QTextDocument, QTextCursor, QTextBlockFormat, QTextCharFormat,
                            QTextFormat, QTextDocumentFragment, QTextDocumentWriter, QFont, QPdfWriter,
                            QPageSize, QPageLayout, QTextListFormat)

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:
    docx = None

PAGE_MARGIN_MM = 20
COVER_TITLE_SIZE = 30
COVER_TOP_MARGIN = 180
CHAPTER_TITLE_SIZE = 20
PARAGRAPH_SPACING = 12
PX_TO_PT = 0.75
PAGE_BREAK_BEFORE = QTextFormat.PageBreakFlag.PageBreak_AlwaysBefore


class ExportError(Exception):
    """Fallo al exportar, con un mensaje apto para mostrar al usuario."""


# ---------------------------------------------------------------------- #
# Construcción del documento
# ---------------------------------------------------------------------- #
def build_document(chapters, info=None) -> QTextDocument:
    """`chapters`: filas con `title` y `content` (HTML). `info`: fila de
    `Database.get_project_info`; si se pasa, se antepone una portada."""
    document = QTextDocument()
    cursor = QTextCursor(document)
    if info is not None:
        _insert_cover(cursor, info)
    # Con portada, el título de la novela es el nivel 1 y los capítulos
    # cuelgan de él (nivel 2); sin portada, los capítulos son el nivel 1.
    chapter_level = 2 if info is not None else 1
    for chapter in chapters:
        _insert_chapter(cursor, chapter, chapter_level)
    return document


def _insert_paragraph(cursor, text, size, *, bold=False, align=Qt.AlignLeft,
                      heading_level=0, top_margin=0, page_break=False):
    """Inserta un bloque de párrafo formateado con la alineación, tamaño y márgenes indicados."""
    if not cursor.document().isEmpty():
        cursor.insertBlock()
    block_format = QTextBlockFormat()
    block_format.setAlignment(align)
    block_format.setHeadingLevel(heading_level)
    block_format.setTopMargin(top_margin)
    block_format.setBottomMargin(PARAGRAPH_SPACING)
    if page_break:
        block_format.setPageBreakPolicy(PAGE_BREAK_BEFORE)
    char_format = QTextCharFormat()
    char_format.setFontPointSize(size)
    char_format.setFontWeight(QFont.Weight.Bold if bold else QFont.Weight.Normal)
    cursor.setBlockFormat(block_format)
    cursor.insertText(text, char_format)


def _insert_cover(cursor, info):
    """Inserta la página de portada en el cursor con los datos generales de la novela."""
    _insert_paragraph(cursor, info["title"] or info["name"], COVER_TITLE_SIZE, bold=True,
                      align=Qt.AlignCenter, heading_level=1, top_margin=COVER_TOP_MARGIN)
    for text, size in ((info["author"], 16), (info["genre"], 12)):
        if text:
            _insert_paragraph(cursor, text, size, align=Qt.AlignCenter, top_margin=16)
    if info["synopsis"]:
        _insert_paragraph(cursor, info["synopsis"], 11, top_margin=48)


def _insert_chapter(cursor, chapter, heading_level):
    """Inserta el título y el cuerpo HTML del capítulo en el documento."""
    _insert_paragraph(cursor, chapter["title"] or "Sin título", CHAPTER_TITLE_SIZE, bold=True,
                      align=Qt.AlignCenter, heading_level=heading_level,
                      page_break=not cursor.document().isEmpty())
    _insert_html_body(cursor, chapter["content"])


def _insert_html_body(cursor, html):
    """Inserta el contenido HTML del capítulo preservando el formato de bloque inicial."""
    body = QTextDocument()
    body.setHtml(html)
    cursor.insertBlock(QTextBlockFormat(), QTextCharFormat())
    start = cursor.position()
    cursor.insertFragment(QTextDocumentFragment(body))
    # Qt descarta el formato (alineación, márgenes) del primer bloque al
    # insertar un fragmento sobre un bloque existente: se restituye a mano.
    first_block = QTextCursor(cursor.document().findBlock(start))
    first_block.setBlockFormat(body.begin().blockFormat())


# ---------------------------------------------------------------------- #
# Escritores por formato
# ---------------------------------------------------------------------- #
def _write_pdf(path, document):
    """Genera un archivo PDF A4 a partir del QTextDocument."""
    writer = QPdfWriter(path)
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageMargins(QMarginsF(PAGE_MARGIN_MM, PAGE_MARGIN_MM, PAGE_MARGIN_MM, PAGE_MARGIN_MM),
                          QPageLayout.Unit.Millimeter)
    document.print_(writer)


def _write_odt(path, document):
    """Escribe el QTextDocument en formato OpenDocument (.odt)."""
    if not QTextDocumentWriter(path, QByteArray(b"ODF")).write(document):
        raise ExportError("No se pudo escribir el archivo .odt.")


def _write_txt(path, document):
    """Vuelca el contenido en texto plano UTF-8."""
    Path(path).write_text(document.toPlainText(), encoding="utf-8")


def _write_markdown(path, document):
    """Convierte los bloques y listas del QTextDocument a Markdown."""
    blocks = []
    block = document.begin()
    while block.isValid():
        content = _markdown_block_content(block)
        if content:
            heading_level = block.blockFormat().headingLevel()
            if heading_level:
                content = f"{'#' * heading_level} {content}"
            text_list = block.textList()
            if text_list is not None:
                list_format = text_list.format()
                indent = "  " * max(0, list_format.indent() - 1)
                if list_format.style() >= 0:
                    content = f"{indent}{text_list.itemNumber(block) + 1}. {content}"
                else:
                    content = f"{indent}- {content}"
            blocks.append(content)
        block = block.next()
    Path(path).write_text("\n\n".join(blocks) + "\n", encoding="utf-8")


def _markdown_block_content(block):
    """Extrae y formatea los fragmentos de un bloque de texto en sintaxis Markdown."""
    parts = []
    fragment = block.begin()
    while not fragment.atEnd():
        text_fragment = fragment.fragment()
        if text_fragment.isValid():
            content = html_escape(text_fragment.text(), quote=False).replace("\u2028", "<br>")
            content = content.replace("\\", "\\\\")
            for marker in ("*", "_", "~", "[", "]"):
                content = content.replace(marker, f"\\{marker}")
            char_format = text_fragment.charFormat()
            if char_format.fontWeight() >= QFont.Weight.Bold:
                content = f"<strong>{content}</strong>"
            if char_format.fontItalic():
                content = f"<em>{content}</em>"
            if char_format.fontUnderline():
                content = f"<u>{content}</u>"
            if char_format.fontStrikeOut():
                content = f"<s>{content}</s>"

            styles = []
            font_families = char_format.fontFamilies()
            if font_families:
                styles.append(f"font-family: '{html_escape(font_families[0], quote=True)}'")
            font_size = char_format.fontPointSize()
            if font_size > 0:
                styles.append(f"font-size: {font_size:g}pt")
            foreground = char_format.foreground()
            if foreground.style() != Qt.BrushStyle.NoBrush:
                color = foreground.color()
                if color.isValid():
                    styles.append(f"color: {color.name()}")
            if styles:
                content = f'<span style="{"; ".join(styles)}">{content}</span>'
            if char_format.isAnchor() and char_format.anchorHref():
                href = html_escape(char_format.anchorHref(), quote=True)
                content = f'<a href="{href}">{content}</a>'
            parts.append(content)
        fragment += 1
    return "".join(parts).strip()


def _write_docx(path, document):
    """Convierte el QTextDocument en un archivo Word .docx con python-docx."""
    if docx is None:
        raise ExportError("Para exportar a .docx hace falta instalar python-docx "
                          "(pip install python-docx).")
    word = docx.Document()
    for style_name in ("Heading 1", "Heading 2"):
        word.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    block = document.begin()
    while block.isValid():
        _add_docx_paragraph(word, block)
        block = block.next()
    word.save(path)


def _add_docx_paragraph(word, block):
    """Añade un párrafo formateado al documento docx."""
    block_format = block.blockFormat()
    level = block_format.headingLevel()
    paragraph = word.add_heading("", level) if level else word.add_paragraph()
    paragraph.alignment = _docx_alignment(block_format.alignment())
    paragraph.paragraph_format.page_break_before = bool(
        block_format.pageBreakPolicy() & PAGE_BREAK_BEFORE)
    if block_format.topMargin() > 0:
        paragraph.paragraph_format.space_before = Pt(block_format.topMargin() * PX_TO_PT)
    if block_format.bottomMargin() > 0:
        paragraph.paragraph_format.space_after = Pt(block_format.bottomMargin() * PX_TO_PT)
    fragments = block.begin()
    while not fragments.atEnd():
        _add_docx_run(paragraph, fragments.fragment())
        fragments += 1


def _add_docx_run(paragraph, fragment):
    """Añade un fragmento de texto con formato de carácter a un párrafo docx."""
    char_format = fragment.charFormat()
    run = paragraph.add_run(fragment.text().replace("\u2028", "\n"))
    run.bold = char_format.fontWeight() >= QFont.Weight.Bold
    run.italic = char_format.fontItalic()
    run.underline = char_format.fontUnderline()
    run.font.strike = char_format.fontStrikeOut()
    font_families = char_format.fontFamilies()
    if font_families:
        run.font.name = font_families[0]
    font_size = char_format.fontPointSize()
    if font_size > 0:
        run.font.size = Pt(font_size)
    foreground = char_format.foreground()
    if foreground.style() != Qt.BrushStyle.NoBrush:
        color = foreground.color()
        if color.isValid():
            run.font.color.rgb = RGBColor(color.red(), color.green(), color.blue())


def _docx_alignment(alignment):
    """Convierte la alineación de Qt en la alineación equivalente de python-docx."""
    if alignment & Qt.AlignHCenter:
        return WD_ALIGN_PARAGRAPH.CENTER
    if alignment & Qt.AlignRight:
        return WD_ALIGN_PARAGRAPH.RIGHT
    if alignment & Qt.AlignJustify:
        return WD_ALIGN_PARAGRAPH.JUSTIFY
    return WD_ALIGN_PARAGRAPH.LEFT


# extensión -> (nombre visible, función de escritura)
FORMATS = {
    "docx": ("Word", _write_docx),
    "pdf": ("PDF", _write_pdf),
    "odt": ("OpenDocument", _write_odt),
    "txt": ("Texto plano", _write_txt),
    "md": ("Markdown", _write_markdown),
}


def export_chapters(path, extension, chapters, info=None):
    """Escribe `chapters` (con portada si se da `info`) en `path`.
    Lanza `ExportError` u `OSError` si no se puede completar."""
    _label, write = FORMATS[extension]
    write(path, build_document(chapters, info))